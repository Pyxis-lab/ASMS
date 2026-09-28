import datetime

from django.contrib.auth.mixins import LoginRequiredMixin, PermissionRequiredMixin
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView, TemplateView, View
from django.urls import reverse_lazy
from django.shortcuts import render
from django.http import HttpResponse, HttpResponseBadRequest
from django.utils.translation import gettext as _
from functools import wraps
from django.contrib import messages
from django.db.models import Q
from django.db import transaction
from django.core.paginator import Paginator
from django.views.decorators.http import require_http_methods
from django.utils.decorators import method_decorator

import openpyxl
from io import BytesIO
import logging

from .models import AfterSalesRecord

logger = logging.getLogger(__name__)

EXCEL_COLUMN_MAP = [
    ('user_name', '用户名称'),
    ('after_sales_type', '售后类型'),
    ('no', 'NO.'),
    ('after_sales_no', '售后编号'),
    ('project_site', '项目现场'),
    ('resolved', '是否解决'),
    ('return_date', '退货时间'),
    ('product_name_model', '货物品名及型号'),
    ('original_ship_date', '发货时间（原件）'),
    ('out_of_warranty', '是否过保'),
    ('sn', 'SN'),
    ('pn', 'PN'),
    ('return_reason_description', '退货原因描述'),
    ('rca', 'RCA'),
    ('after_sales_cost', '售后成本'),
    ('charge_details', '收费明细'),
    (None, 'ShippingInformation'),  # blank column
    ('ship_date', '寄货时间'),
    ('customer_return_required', '是否需客户寄回'),
    ('shipped_item_name', '寄货品名'),
    ('r_sn', 'R_SN'),
    ('r_pn', 'R_PN'),
    ('equipment_no', '设备编号'),
    ('outbound_order_no', '出库单号'),
    ('shipping_info', '寄货信息'),
    ('tracking_no', '运单号'),
    ('source_warehouse', '拿货仓库'),
]



def parse_boolean_value(value):
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    str_value = str(value).strip()
    if str_value in ('是', 'Y', 'Yes', 'YES', 'true', 'True', 'TRUE', '1'):
        return True
    return False


def parse_datetime_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return value.strftime('%Y/%m/%d')
    if isinstance(value, datetime.date):
        return value.strftime('%Y/%m/%d')
    if isinstance(value, float):
        try:
            dt = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=value)
            return dt.strftime('%Y/%m/%d')
        except (OverflowError, ValueError):
            pass
    str_value = str(value).strip()
    datetime_formats = [
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%Y-%m-%d',
        '%Y/%m/%d %H:%M:%S',
        '%Y/%m/%d %H:%M',
        '%Y/%m/%d',
        '%d-%m-%Y ',
        '%d/%m/%Y ',
    ]
    for fmt in datetime_formats:
        try:
            dt = datetime.datetime.strptime(str_value, fmt)
            return dt.strftime('%Y/%m/%d')
        except ValueError:
            continue
    return str_value


def validate_row(row_data, row_num):
    errors = []
    after_sales_no = row_data.get('after_sales_no', '').strip()
    if not after_sales_no:
        errors.append(f"第{row_num}行: 售后编号不能为空")
    return errors


class AfterSalesRecordListView(LoginRequiredMixin, ListView):
    model = AfterSalesRecord
    template_name = 'product/aftersales_list.html'
    context_object_name = 'records'
    paginate_by = 50
    login_url = '/login/'

    def get_queryset(self):
        queryset = super().get_queryset()
        search_query = self.request.GET.get('search', '')
        after_sales_type = self.request.GET.get('after_sales_type', '')
        resolved = self.request.GET.get('resolved', '')
        out_of_warranty = self.request.GET.get('out_of_warranty', '')

        if search_query:
            queryset = queryset.filter(
                Q(after_sales_no__icontains=search_query) |
                Q(sn__icontains=search_query) |
                Q(pn__icontains=search_query) |
                Q(user_name__icontains=search_query) |
                Q(product_name_model__icontains=search_query)
            )

        if after_sales_type:
            queryset = queryset.filter(after_sales_type=after_sales_type)

        if resolved != '':
            queryset = queryset.filter(resolved=resolved == 'True')

        if out_of_warranty != '':
            queryset = queryset.filter(out_of_warranty=out_of_warranty == 'True')

        return queryset.order_by('-return_date')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search_query'] = self.request.GET.get('search', '')
        context['selected_type'] = self.request.GET.get('after_sales_type', '')
        context['selected_resolved'] = self.request.GET.get('resolved', '')
        context['selected_out_of_warranty'] = self.request.GET.get('out_of_warranty', '')
        context['after_sales_types'] = AfterSalesRecord.objects.values_list(
            'after_sales_type', flat=True
        ).distinct().exclude(after_sales_type='')
        return context


class AfterSalesRecordDetailView(LoginRequiredMixin, DetailView):
    model = AfterSalesRecord
    template_name = 'product/aftersales_detail.html'
    context_object_name = 'record'
    login_url = '/login/'


class AfterSalesRecordCreateView(LoginRequiredMixin, PermissionRequiredMixin, CreateView):
    model = AfterSalesRecord
    permission_required = 'product.add_aftersalesrecord'
    template_name = 'product/aftersales_form.html'
    context_object_name = 'record'
    fields = [
        'after_sales_no', 'no', 'user_name', 'after_sales_type', 'project_site',
        'product_name_model', 'sn', 'pn', 'equipment_no',
        'resolved', 'out_of_warranty',
        'return_date', 'original_ship_date', 'ship_date',
        'return_reason_description', 'rca',
        'shipped_item_name', 'r_sn', 'r_pn',
        'customer_return_required', 'outbound_order_no', 'tracking_no', 'source_warehouse',
        'shipping_info', 'after_sales_cost', 'charge_details',
    ]
    success_url = reverse_lazy('product:aftersales_list')
    login_url = '/login/'

    def form_valid(self, form):
        messages.success(self.request, '售后记录创建成功')
        return super().form_valid(form)

    def form_invalid(self, form):
        messages.error(self.request, '表单验证失败，请检查输入')
        return super().form_invalid(form)


class AfterSalesRecordUpdateView(LoginRequiredMixin, PermissionRequiredMixin, UpdateView):
    model = AfterSalesRecord
    permission_required = 'product.change_aftersalesrecord'
    template_name = 'product/aftersales_form.html'
    context_object_name = 'record'
    fields = [
        'after_sales_no', 'no', 'user_name', 'after_sales_type', 'project_site',
        'product_name_model', 'sn', 'pn', 'equipment_no',
        'resolved', 'out_of_warranty',
        'return_date', 'original_ship_date', 'ship_date',
        'return_reason_description', 'rca',
        'shipped_item_name', 'r_sn', 'r_pn',
        'customer_return_required', 'outbound_order_no', 'tracking_no', 'source_warehouse',
        'shipping_info', 'after_sales_cost', 'charge_details',
    ]
    success_url = reverse_lazy('product:aftersales_list')
    login_url = '/login/'

    def form_valid(self, form):
        messages.success(self.request, '售后记录更新成功')
        return super().form_valid(form)

    def form_invalid(self, form):
        messages.error(self.request, '表单验证失败，请检查输入')
        return super().form_invalid(form)


class AfterSalesRecordDeleteView(LoginRequiredMixin, PermissionRequiredMixin, DeleteView):
    model = AfterSalesRecord
    template_name = 'product/aftersales_confirm_delete.html'
    permission_required = 'product.delete_aftersalesrecord'
    success_url = reverse_lazy('product:aftersales_list')
    login_url = '/login/'

    def form_valid(self, form):
        messages.success(self.request, '售后记录删除成功')
        return super().form_valid(form)


class AfterSalesRecordImportView(LoginRequiredMixin, PermissionRequiredMixin, TemplateView):
    template_name = 'product/aftersales_import.html'
    login_url = '/login/'
    permission_required = 'product.add_aftersalesrecord'

    def get(self, request, *args, **kwargs):
        return render(request, self.template_name)

    def post(self, request, *args, **kwargs):
        if request.POST.get('action') == 'confirm':
            return self.process_import(request)

        if 'excel_file' not in request.FILES:
            return render(request, self.template_name, {
                'error_message': '请选择要上传的Excel文件'
            })

        excel_file = request.FILES['excel_file']
        if not excel_file.name.endswith('.xlsx'):
            return render(request, self.template_name, {
                'error_message': '只支持.xlsx格式的Excel文件'
            })

        update_existing = request.POST.get('update_existing') == 'on'

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            ws = wb.active

            preview_data = []
            error_list = []
            seen_numbers = {}

            for row_num in range(2, ws.max_row + 1):
                row_data = {}
                for col_idx, (field_name, _) in enumerate(EXCEL_COLUMN_MAP):
                    if field_name is None:
                        continue
                    cell_value = ws.cell(row=row_num, column=col_idx + 1).value
                    if cell_value is not None:
                        if isinstance(cell_value, float) and cell_value.is_integer():
                            cell_value = int(cell_value)
                        row_data[field_name] = str(cell_value).strip()

                # Skip blank rows (formatted but empty rows at the end of a sheet)
                if not any(row_data.values()):
                    continue

                row_errors = validate_row(row_data, row_num)
                if row_errors:
                    error_list.append({'row': row_num, 'message': row_errors[0]})
                    continue

                # after_sales_no is unique: a repeat inside the file would fail the whole import
                after_sales_no = row_data['after_sales_no']
                if after_sales_no in seen_numbers:
                    error_list.append({
                        'row': row_num,
                        'message': f"第{row_num}行: 售后编号 '{after_sales_no}' 与第{seen_numbers[after_sales_no]}行重复",
                    })
                    continue
                seen_numbers[after_sales_no] = row_num

                if 'resolved' in row_data:
                    row_data['resolved'] = parse_boolean_value(row_data['resolved'])
                if 'out_of_warranty' in row_data:
                    row_data['out_of_warranty'] = parse_boolean_value(row_data['out_of_warranty'])
                if 'customer_return_required' in row_data:
                    row_data['customer_return_required'] = parse_boolean_value(row_data['customer_return_required'])

                datetime_fields = ['return_date', 'original_ship_date', 'ship_date']
                for field in datetime_fields:
                    if field in row_data:
                        row_data[field] = parse_datetime_value(row_data[field])

                action_type = 'create'
                original_data = None
                if not update_existing and AfterSalesRecord.objects.filter(after_sales_no=after_sales_no).exists():
                    error_list.append({
                        'row': row_num,
                        'message': f"第{row_num}行: 售后编号 '{after_sales_no}' 已存在（未勾选更新已有记录）",
                    })
                    continue
                if update_existing:
                    try:
                        existing_record = AfterSalesRecord.objects.get(after_sales_no=after_sales_no)
                        action_type = 'update'
                        original_data = {}
                        for field_name, _ in EXCEL_COLUMN_MAP:
                            if field_name and hasattr(existing_record, field_name):
                                value = getattr(existing_record, field_name)
                                if value is not None:
                                    original_data[field_name] = str(value)
                    except AfterSalesRecord.DoesNotExist:
                        action_type = 'create'

                changed_fields = {}
                if action_type == 'update' and original_data:
                    for field_name, new_value in row_data.items():
                        if field_name in original_data and str(new_value) != original_data[field_name]:
                            changed_fields[field_name] = {
                                'original': original_data[field_name],
                                'new': new_value,
                            }

                preview_data.append({
                    'row_num': row_num,
                    'action': action_type,
                    'data': row_data,
                    'changed_fields': changed_fields,
                })

            create_count = sum(1 for p in preview_data if p['action'] == 'create')
            update_count = sum(1 for p in preview_data if p['action'] == 'update')

            request.session['import_preview'] = {
                'preview_data': preview_data,
                'update_existing': update_existing,
                'file_name': excel_file.name,
            }

            return render(request, self.template_name, {
                'preview_mode': True,
                'preview_data': preview_data,
                'create_count': create_count,
                'update_count': update_count,
                'error_count': len(error_list),
                'file_name': excel_file.name,
                'errors': error_list,
            })

        except Exception as e:
            logger.error(f'Excel导入失败: {str(e)}', exc_info=True)
            return render(request, self.template_name, {
                'error_message': f'导入失败：{str(e)}'
            })

    def process_import(self, request):
        import_preview = request.session.get('import_preview')
        if not import_preview:
            return render(request, self.template_name, {
                'error_message': '请先上传文件并预览'
            })

        preview_data = import_preview['preview_data']
        update_existing = import_preview.get('update_existing', True)

        success_count = 0
        updated_count = 0
        failed_count = 0
        error_list = []

        batch_size = 100
        records_to_create = []
        records_to_update = []

        try:
            with transaction.atomic():
                for item in preview_data:
                    row_data = item['data']
                    row_num = item['row_num']

                    try:
                        if item['action'] == 'update':
                            existing_record = AfterSalesRecord.objects.get(after_sales_no=row_data['after_sales_no'])
                            for key, value in row_data.items():
                                setattr(existing_record, key, value)
                            records_to_update.append(existing_record)
                            updated_count += 1
                        else:
                            records_to_create.append(AfterSalesRecord(**row_data))
                            success_count += 1

                    except AfterSalesRecord.DoesNotExist:
                        records_to_create.append(AfterSalesRecord(**row_data))
                        success_count += 1
                    except Exception as e:
                        error_list.append({'row': row_num, 'message': str(e)})
                        failed_count += 1

                    if len(records_to_create) >= batch_size:
                        AfterSalesRecord.objects.bulk_create(records_to_create)
                        records_to_create = []

                    if len(records_to_update) >= batch_size:
                        AfterSalesRecord.objects.bulk_update(
                            records_to_update,
                            fields=[f for f, _ in EXCEL_COLUMN_MAP if f and f != 'after_sales_no']
                        )
                        records_to_update = []

                if records_to_create:
                    AfterSalesRecord.objects.bulk_create(records_to_create)

                if records_to_update:
                    AfterSalesRecord.objects.bulk_update(
                        records_to_update,
                        fields=[f for f, _ in EXCEL_COLUMN_MAP if f and f != 'after_sales_no']
                    )

            import_result = {
                'total': len(preview_data),
                'success': success_count,
                'updated': updated_count,
                'failed': failed_count,
                'errors': error_list[:50],
            }

            if failed_count == 0:
                messages.success(request, f'成功导入 {success_count} 条记录，更新 {updated_count} 条记录')
            else:
                messages.warning(request, f'导入完成：成功 {success_count} 条，更新 {updated_count} 条，失败 {failed_count} 条')

            del request.session['import_preview']

            return render(request, self.template_name, {'import_result': import_result})

        except Exception as e:
            logger.error(f'Excel导入失败: {str(e)}', exc_info=True)
            del request.session['import_preview']
            return render(request, self.template_name, {
                'error_message': f'导入失败（已回滚，未导入任何记录）：{str(e)}'
            })


class AfterSalesRecordExportView(LoginRequiredMixin, PermissionRequiredMixin, View):
    login_url = '/login/'
    permission_required = 'product.view_aftersalesrecord'

    def get(self, request, *args, **kwargs):
        try:
            records = AfterSalesRecord.objects.all().order_by('-return_date')

            output = BytesIO()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = '售后记录'

            header_row = []
            for _, label in EXCEL_COLUMN_MAP:
                header_row.append(label if label else '')
            ws.append(header_row)

            for record in records:
                row_data = []
                for field_name, _ in EXCEL_COLUMN_MAP:
                    if field_name is None:
                        row_data.append('')
                    else:
                        value = getattr(record, field_name, '')
                        if isinstance(value, bool):
                            row_data.append('是' if value else '否')
                        elif value is None:
                            row_data.append('')
                        else:
                            row_data.append(str(value))
                ws.append(row_data)

            for col in ws.columns:
                max_length = 0
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[col[0].column_letter].width = adjusted_width

            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output,
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename=售后记录.xlsx'
            return response

        except Exception as e:
            logger.error(f'Excel导出失败: {str(e)}', exc_info=True)
            messages.error(request, f'导出失败：{str(e)}')
            return HttpResponseBadRequest(f'导出失败：{str(e)}')